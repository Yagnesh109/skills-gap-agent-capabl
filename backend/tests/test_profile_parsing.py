from io import BytesIO
from unittest.mock import patch

import pytest
from docx import Document
from fastapi.testclient import TestClient

from main import app
from profile_parsing.schemas import EMPTY_PROFILE
from profile_parsing.document_extractor import extract_text_from_file
from profile_parsing.pdf_extractor import extract_text_from_pdf

client = TestClient(app)


def test_empty_profile_has_expected_structure():
    assert EMPTY_PROFILE["personal_info"]["name"] == ""
    assert EMPTY_PROFILE["education"] == []
    assert EMPTY_PROFILE["skills"] == []


def test_invalid_pdf_bytes_raise_value_error():
    with pytest.raises(ValueError):
        extract_text_from_pdf(b"this is not a pdf")


def test_docx_text_is_extracted():
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Python Developer")
    document_bytes = BytesIO()
    document.save(document_bytes)

    extracted_text = extract_text_from_file(document_bytes.getvalue(), "resume.docx")

    assert "Jane Doe" in extracted_text
    assert "Python Developer" in extracted_text


def test_txt_text_is_extracted():
    resume_text = b"Alex Morgan\nBackend Developer\nPython, FastAPI"

    extracted_text = extract_text_from_file(resume_text, "resume.txt")

    assert extracted_text == "Alex Morgan\nBackend Developer\nPython, FastAPI"


def test_parse_profile_text_empty_error():
    response = client.post("/api/profile/parse-text", json={"text": "   "})
    assert response.status_code == 400
    assert response.json()["detail"] == "Resume text cannot be empty."


@patch("main.parse_resume_with_gemini")
def test_parse_profile_text_success(mock_parse):
    mock_parse.return_value = {
        **EMPTY_PROFILE,
        "personal_info": {"name": "John Doe", "email": "john@example.com", "phone": "", "location": ""},
        "skills": ["Python", "FastAPI"],
    }
    response = client.post("/api/profile/parse-text", json={"text": "John Doe, Python Developer with FastAPI"})
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["profile"]["personal_info"]["name"] == "John Doe"
    assert data["profile"]["skills"] == ["Python", "FastAPI"]
    mock_parse.assert_called_once_with("John Doe, Python Developer with FastAPI")


@patch("main.parse_resume_image_with_gemini")
def test_image_resume_parsing(mock_image_parse):
    mock_image_parse.return_value = {
        **EMPTY_PROFILE,
        "personal_info": {"name": "Sarah Connor", "email": "", "phone": "", "location": ""},
        "skills": ["Cybersecurity"],
    }
    from fastapi.testclient import TestClient
    from main import app

    client = TestClient(app)
    response = client.post(
        "/api/profile/parse",
        files={"file": ("resume.png", b"fake_png_bytes", "image/png")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["profile"]["personal_info"]["name"] == "Sarah Connor"
    assert data["profile"]["skills"] == ["Cybersecurity"]


@patch("google.generativeai.GenerativeModel")
def test_gemini_fallback_on_429_retry(mock_model_cls):
    from profile_parsing.gemini_client import parse_resume_with_gemini

    class Mock429Model:
        def generate_content(self, *args, **kwargs):
            raise Exception("429 You exceeded your current quota, limit: 5, model: gemini-3.6-flash")

    mock_model_cls.return_value = Mock429Model()

    profile = parse_resume_with_gemini("Alice Smith Python Developer")
    assert "Alice Smith" in profile["personal_info"]["name"]
    assert "Python" in profile["skills"]





